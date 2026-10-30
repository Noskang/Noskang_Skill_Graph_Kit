"""Reading, expansion, validation and export on the demo sample."""
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from helpers import SAMPLE, Temp, codes, error_codes, find, sha
from skillgraph import demo
from skillgraph.core import SkillDataError
from skillgraph.keyframes import split_keyframe
from skillgraph.pipeline import build, run
from skillgraph.views import dictionary_view, expand_view, graph_view


def check(t, product=True):
    return run(t.cfg.documents(), t.cfg, product)


class SampleTests(unittest.TestCase):
    def test_demo_is_clean(self):
        with Temp() as t:
            r = check(t)
            self.assertEqual((r['errors'], r['warnings']), ([], []))
            self.assertEqual((len(r['logic']['skills']), len(r['logic']['buffs'])), (9, 2))

    def test_committed_sample_is_reproducible(self):
        # The committed workbooks, dictionary and clips are exactly what the generator writes.
        with tempfile.TemporaryDirectory() as tmp:
            demo.make(tmp)
            fresh = sorted(p.relative_to(tmp).as_posix() for p in Path(tmp).rglob('*') if p.is_file())
            committed = sorted(p.relative_to(SAMPLE).as_posix() for p in SAMPLE.rglob('*') if p.is_file() and p.suffix != '.meta')
            self.assertEqual(fresh, committed)
            for name in fresh:
                self.assertEqual(sha(Path(tmp) / name), sha(SAMPLE / name), name)

    def test_same_input_same_fingerprint(self):
        with Temp() as t:
            a = build(t.cfg)['manifest']
            b = build(t.cfg)['manifest']
            self.assertEqual(a['logicSha256'], b['logicSha256'])
            self.assertEqual(a['presentationSha256'], b['presentationSha256'])

    def test_expansion_layers(self):
        with Temp() as t:
            r = check(t)
            mods = r['mods']
            self.assertEqual(mods[2200203]['chain'], [2200001, 2200201, 2200203])
            self.assertIn(100, mods[2200202]['nodes'], "tier 2 adds a node with '+'")
            self.assertNotIn(12, mods[2200203]['nodes'], "tier 3 removes the sound with '-'")
            self.assertEqual(mods[2200202]['links'][(11, 'AngleStep')]['_from'], 2200202, 'slot override in a child layer')
            self.assertEqual(r['vals'][2300202]['Bolt.Damage'], (30, 2300202))
            self.assertEqual(r['vals'][2300202]['Bolt.Speed'], (14.0, 2300201), 'value inherited from the parent bundle')
            self.assertEqual(r['reach'][2100205], [2200205, 2200215, 2200216])

    def test_time_values_snap_to_grid(self):
        with Temp(rows=lambda rows: find(rows, 'ValueData', BundleKey=2300201, Name='Bolt.Release').update(Value=0.31)) as t:
            r = check(t)
            self.assertIn('OffGrid', codes(r))
            entries = dict(next(v for v in r['logic']['values'] if v['key'] == 2300201)['entries'])
            idx = r['names'].index('Bolt.Release')
            self.assertEqual(entries[idx], 0.3, 'exported time is snapped to the 0.025 s grid')

    def test_keyframe_reference(self):
        with Temp() as t:
            r = check(t)
            slash = next(b for b in r['logic']['bundles'] if b['key'] == 2200101)
            self.assertEqual(next(n['at'] for n in slash['nodes'] if n['id'] == 11), 0.25, '@10:Swing = frame 15 at 60 fps')
            pres = next(b for b in r['presentation']['bundles'] if b['key'] == 2200101)
            self.assertEqual(next(n['at'] for n in pres['nodes'] if n['id'] == 12), 0.25)
        self.assertEqual(split_keyframe('Swing+0.1', {'Swing'}), ('Swing', 0.1))
        self.assertEqual(split_keyframe('Hit_2', {'Hit_2'}), ('Hit_2', 0.0))

        def set_ref(value):
            return lambda rows: find(rows, 'ModuleNodeData', BundleKey=2200101, NodeId=11).update(EventValue=value)
        cases = [('@10:Swing+0.05', None, 0.3), ('@10:Swing-0.1', None, 0.15), ('@10:Start-0.1', 'KeyframeOffsetNegative', None),
                 ('@10:Nope', 'KeyframeMissing', None), ('@1:Swing', 'KeyframeMotion', None)]
        for value, code, at in cases:
            with Temp(rows=set_ref(value)) as t:
                r = check(t)
                if code:
                    self.assertIn(code, error_codes(r), value)
                else:
                    slash = next(b for b in r['logic']['bundles'] if b['key'] == 2200101)
                    self.assertEqual(next(n['at'] for n in slash['nodes'] if n['id'] == 11), at, value)

    def test_value_name_rules(self):
        def edit(rows):
            rows['ValueData'].append(dict(BundleKey=2300101, Name='hit.damage', Value=1))
            rows['ValueData'].append(dict(BundleKey=2300101, Name='Hit.DAMAGE', Value=1))
        with Temp(rows=edit) as t:
            self.assertTrue({'ValueName', 'ValueNameCase'} <= set(error_codes(check(t))))

    def test_key_band(self):
        with Temp(rows=lambda rows: find(rows, 'SkillMetaData', Key=2100101).update(ValueBundleKey=2200101)) as t:
            self.assertIn('KeyBand', error_codes(check(t)))

    def test_inheritance_cycle_and_depth(self):
        with Temp(rows=lambda rows: find(rows, 'ModuleBundleData', Key=2200001).update(ParentKey=2200101)) as t:
            self.assertIn('InheritanceCycle', error_codes(check(t)))
        with Temp(rows=lambda rows: find(rows, 'ModuleBundleData', Key=2200203).update(ParentKey=2200202)) as t:
            self.assertIn('InheritanceDepth', error_codes(check(t)), 'more than three layers')

    def test_missing_value_name(self):
        with Temp(rows=lambda rows: rows['ValueData'].remove(find(rows, 'ValueData', BundleKey=2300101, Name='Hit.Radius'))) as t:
            r = check(t)
            self.assertIn('MissingValueName', error_codes(r))
            issue = next(x for x in r['errors'] if x.code == 'MissingValueName')
            self.assertEqual((issue.at['bundle'], issue.at['node'], issue.at['slot']), (2200002, 2, 'Radius'))

    def test_enum_number_is_rejected(self):
        with Temp(rows=lambda rows: find(rows, 'ModuleBundleData', Key=2200001).update(Kind=0)) as t:
            self.assertIn('EnumName', error_codes(check(t)))

    def test_fixed_balance_warning(self):
        with Temp(rows=lambda rows: find(rows, 'ValueLinkData', BundleKey=2200002, NodeId=3, Slot='Amount').update(Source='Fixed', Value=5)) as t:
            r = check(t)
            self.assertIn('FixedBalance', [x.code for x in r['warnings']])

    def test_presentation_event(self):
        def edit(rows):
            rows['ModuleNodeData'].append(demo.node(2200221, 2, 'StatusRemove', 'PlaySound', line='Presentation'))
            rows['ValueLinkData'].append(demo.fix(2200221, 2, 'PresentationKey', 2500106))
        with Temp(rows=edit) as t:
            self.assertIn('PresentationEvent', error_codes(check(t)))

    def test_companion_rule(self):
        with Temp(rows=lambda rows: find(rows, 'ModuleNodeData', BundleKey=2200214, NodeId=2).update(Module='HitWindow')) as t:
            r = check(t)
            self.assertIn('CompanionMissing', error_codes(r))
        with Temp(settings={'rules': {'companions': []}}, rows=lambda rows: find(rows, 'ModuleNodeData', BundleKey=2200214, NodeId=2).update(Module='HitWindow')) as t:
            self.assertNotIn('CompanionMissing', error_codes(check(t)), 'the rule comes from settings')

    def test_product_only_rules(self):
        def edit(rows):
            rows['ModuleNodeData'].append(demo.node(2200002, 4, 'Hit', 'Knockback', order=30))
            rows['ValueLinkData'].append(demo.fix(2200002, 4, 'Force', 1.0))
            find(rows, 'ValueLinkData', BundleKey=2200215, NodeId=4, Slot='Kind').update(Value='Pierce')
        with Temp(rows=edit) as t:
            product = error_codes(check(t, True))
            self.assertTrue({'UnimplementedModule', 'UnsupportedValue'} <= set(product))
            self.assertEqual(error_codes(check(t, False)), [], 'sample mode checks shape only')

    def test_order_tie(self):
        with Temp(rows=lambda rows: find(rows, 'ModuleNodeData', BundleKey=2200002, NodeId=2).update(Order=10)) as t:
            self.assertIn('OrderTie', error_codes(check(t)))

    def test_required_slot(self):
        with Temp(rows=lambda rows: rows['ValueLinkData'].remove(find(rows, 'ValueLinkData', BundleKey=2200001, NodeId=1, Slot='Duration'))) as t:
            self.assertIn('RequiredSlot', error_codes(check(t)))

    def test_spawn_must_target_effect(self):
        with Temp(rows=lambda rows: find(rows, 'ValueLinkData', BundleKey=2200101, NodeId=11, Slot='BundleKey').update(Value=2200111)) as t:
            self.assertIn('SpawnNotEffect', error_codes(check(t)))

    def test_meta_bundle_kind(self):
        with Temp(rows=lambda rows: find(rows, 'SkillMetaData', Key=2100101).update(ModuleBundleKey=2200002)) as t:
            self.assertIn('MetaBundleKind', error_codes(check(t)))

    def test_exit_events_need_a_check_node(self):
        with Temp(rows=lambda rows: find(rows, 'ModuleNodeData', BundleKey=2200205, NodeId=11).update(FromNode=1)) as t:
            self.assertIn('FromNodeNotCheck', error_codes(check(t)))

    def test_practice_owner_is_rejected_in_product(self):
        def edit(rows):
            rows['ValueBundleData'].append(dict(Key=2309801, ParentKey=None, Desc='practice'))
        with Temp(rows=edit) as t:
            self.assertIn('PracticeKey', error_codes(check(t)))
            self.assertNotIn('PracticeKey', error_codes(run(t.cfg.documents(), t.cfg, True, practice=True)))

    def test_header_and_sheet_rules(self):
        with Temp() as t:
            path = t.cfg.data_dir / 'Skills.xlsx'
            wb = openpyxl.load_workbook(path)
            wb['Values']['C4'] = 'Amount'
            wb.create_sheet('Scratch')
            wb.create_sheet('!Notes')
            wb['Values']['B9'] = '=A1'
            wb['Values']['G9'] = 'stray'
            wb.save(path)
            r = check(t)
            self.assertTrue({'HeaderMismatch', 'SheetsMismatch', 'FormulaNotAllowed', 'ExtraCell'} <= set(error_codes(r)))
            self.assertFalse([x for x in r['errors'] if '!Notes' in str(x)], '"!" sheets are notes')

    def test_unknown_dictionary_format(self):
        with Temp() as t:
            p = t.cfg.dictionary_path
            d = json.loads(p.read_text(encoding='utf-8'))
            d['format'] = 'Something else'
            p.write_text(json.dumps(d), encoding='utf-8')
            with self.assertRaises(SkillDataError):
                check(t)

    def test_messages_follow_language(self):
        def edit(rows):
            find(rows, 'ModuleNodeData', BundleKey=2200002, NodeId=2).update(Order=10)
        with Temp(rows=edit, lang='ko') as t:
            r = check(t)
            tie = next(x for x in r['errors'] if x.code == 'OrderTie')
            self.assertIn('순서가 같습니다', tie.message)
            self.assertIn('묶음 2200002', tie.where)
            self.assertIn('기술 메타', openpyxl.load_workbook(t.cfg.data_dir / 'Skills.xlsx').sheetnames, 'sheet names follow the language')

    def test_build_raises_on_errors(self):
        with Temp(rows=lambda rows: find(rows, 'ModuleNodeData', BundleKey=2200002, NodeId=2).update(Order=10)) as t:
            with self.assertRaises(SkillDataError):
                build(t.cfg)

    def test_plugins(self):
        plugin = (
            'def no_big_damage(ctx):\n'
            '    for k, v in ctx["vals"].items():\n'
            '        for name, (value, _) in v.items():\n'
            '            if name.endswith(".Damage") and value > 50:\n'
            '                ctx["report"].w("BigDamage", dict(values=k, name=name), "UnusedValue", name=name)\n'
            'def stamp(logic, pres, ctx):\n'
            '    logic["host"] = "demo"\n'
        )
        with Temp() as t:
            (t.dir / 'host_rules.py').write_text(plugin, encoding='utf-8')
            t.cfg.data['plugins'] = {'paths': ['.'], 'validate': ['host_rules:no_big_damage'], 'export': 'host_rules:stamp'}
            r = check(t)
            self.assertEqual([x.code for x in r['warnings']], ['BigDamage'])
            self.assertEqual(r['logic']['host'], 'demo')

    def test_generated_views(self):
        with Temp() as t:
            res = dictionary_view(t.cfg)
            self.assertEqual(res['modules'], len(demo.MODULES))
            wb = openpyxl.load_workbook(res['file'])
            self.assertIn('Module dictionary', wb.sheetnames)
            before = sha(t.cfg.data_dir / 'Skills.xlsx')
            res = expand_view(t.cfg)
            self.assertGreater(res['rows'], 50)
            wb = openpyxl.load_workbook(t.cfg.data_dir / 'Skills.xlsx')
            self.assertIn('!Generated_Expanded', wb.sheetnames)
            self.assertNotEqual(before, sha(t.cfg.data_dir / 'Skills.xlsx'))
            self.assertEqual(check(t)['errors'], [], 'the generated sheet is not read as data')
            self.assertTrue((t.cfg.backup_dir / 'Skills.xlsx').is_file(), 'one previous copy')
            out = t.dir / 'view.json'
            res = graph_view(t.cfg, out)
            model = json.loads(out.read_text(encoding='utf-8'))
            slash = next(b for b in model['bundles'] if b['key'] == 2200101)
            lock = next(n for n in slash['nodes'] if n['module'] == 'CastLock')
            self.assertEqual((lock['definedIn'], lock['slots'][0]['source'], lock['slots'][0]['resolved'][0]['value']), (2200001, 'ref', 0.5))


if __name__ == '__main__':
    unittest.main()
