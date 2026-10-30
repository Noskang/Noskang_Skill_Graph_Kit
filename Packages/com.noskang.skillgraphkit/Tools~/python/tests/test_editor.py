"""Editor backend: operations, model, Excel round trip, locks, practice workspace, aux file, columns, serve."""
import io
import json
import unittest

import openpyxl

from helpers import Temp, sha
from skillgraph import editor as se
from skillgraph.pipeline import build


def rows(loaded, name):
    t = loaded.tables[name]
    return [tuple(r.get(c['name']) for c in t['table']['columns']) for r in t['rows']]


def node(model, bundle, nid):
    b = next(b for b in model['bundles'] if b['key'] == bundle)
    return next((n for n in b['nodes'] if n['id'] == nid), None), b


def errors(model):
    return [i for i in model['issues'] if i['level'] == 'error']


class EditorTests(unittest.TestCase):
    def test_model_without_edits(self):
        with Temp() as t:
            m = se.Backend(t.cfg).handle(dict(cmd='model'))['model']
            self.assertEqual((len(m['owners']), len(m['bundles'])), (11, 19))
            self.assertEqual(errors(m), [])
            self.assertEqual(len(m['display']['families']), 6)
            lock = next(x for x in m['dictionary']['modules'] if x['name'] == 'CastLock')
            self.assertEqual((lock['label'], lock['slots'][0]['description']), ('Cast lock', 'How long the cast lasts'), 'display text merged in')
            t3 = next(o for o in m['owners'] if o['key'] == 2100203)
            self.assertEqual(t3['stage']['family'], [2100201, 2100202, 2100203])
            self.assertEqual(t3['chain'], [2200001, 2200201, 2200203])
            cry = next(o for o in m['owners'] if o['key'] == 2100103)
            self.assertEqual(cry['buffs'], [2110101])
            b = next(b for b in m['bundles'] if b['key'] == 2200201)
            self.assertEqual(b['usedBy'], [2100201, 2100202, 2100203], 'everyone that inherits the bundle')
            self.assertEqual(next(b for b in m['bundles'] if b['key'] == 2200003)['prefixes'], ['Bolt'], 'object names to suggest for new links')
            slash = next(b for b in m['bundles'] if b['key'] == 2200101)
            self.assertEqual([(k['name'], k['seconds']) for k in slash['keyframes'][0]['keys']][:3], [('Start', 0.0), ('Windup', 0.1), ('Swing', 0.25)])
            self.assertEqual(m['settings']['layers'], ['Base', 'Owner skill', 'Tier'])

    def test_value_edit_round_trip_and_revert_restores_fingerprint(self):
        with Temp() as t:
            before = build(t.cfg)['manifest']['logicSha256']
            bk = se.Backend(t.cfg)
            first = se.Loaded(se.Workspace(t.cfg))
            original = {k: rows(first, k) for k in first.tables}
            op = [dict(op='setValue', values=2300201, name='Bolt.Damage', value=26)]
            res = bk.handle(dict(cmd='changes', ops=op))
            self.assertEqual([(c['kind'], c['sheet'], c['column'], c['old'], c['new']) for c in res['changes']], [('change', 'Values', 'Value', 25, 26)])
            res = bk.handle(dict(cmd='save', ops=op, base=first.base))
            self.assertTrue(res['ok'], res)
            again = se.Loaded(se.Workspace(t.cfg))
            changed = [(k, i) for k in again.tables for i, (x, y) in enumerate(zip(rows(again, k), original[k])) if x != y]
            self.assertEqual(len(changed), 1, changed)
            self.assertTrue((t.cfg.backup_dir / 'Skills.xlsx').is_file(), 'one previous copy')
            self.assertNotEqual(build(t.cfg)['manifest']['logicSha256'], before)
            res = bk.handle(dict(cmd='save', ops=[dict(op='setValue', values=2300201, name='Bolt.Damage', value=25)], base=again.base))
            self.assertTrue(res['ok'], res)
            self.assertEqual(build(t.cfg)['manifest']['logicSha256'], before, 'reverting restores the data fingerprint')

    def test_child_layer_add_remove_revert(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            ops = [dict(op='addNode', bundle=2200202, line='Presentation', event='CastStart', module='PlaySound',
                        links=[dict(slot='PresentationKey', source='Fixed', value=2500202)])]
            m = bk.handle(dict(cmd='model', ops=ops))['model']
            new_id = m['ops'][0]['node']
            self.assertGreaterEqual(new_id, 100, 'tier layer ids start at 100')
            n, _ = node(m, 2200202, new_id)
            self.assertTrue(n['added'] and n['definedIn'] == 2200202)
            self.assertIsNone(node(m, 2200201, new_id)[0], 'not in the parent')
            e, _ = se.apply_ops(bk.loaded('product'), ops)
            self.assertEqual(e.find('ModuleNodeData', BundleKey=2200202, NodeId=new_id)['Override'], '+')
            m2 = bk.handle(dict(cmd='model', ops=[dict(op='removeNode', bundle=2200202, node=11)]))['model']
            self.assertIsNone(node(m2, 2200202, 11)[0])
            self.assertIn(11, [x['id'] for x in node(m2, 2200202, 1)[1]['removed']])
            m3 = bk.handle(dict(cmd='model', ops=[dict(op='removeNode', bundle=2200202, node=11), dict(op='revert', bundle=2200202, node=11)]))['model']
            self.assertIsNotNone(node(m3, 2200202, 11)[0])
            with self.assertRaises(se.OpError):
                se.apply_ops(bk.loaded('product'), [dict(op='moveNode', bundle=2200202, node=11, event='CastStart')])

    def test_override_slot_and_apply_to_parent(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            ops = [dict(op='setLink', bundle=2200203, node=11, slot='Count', source='Fixed', value=2)]
            m1 = bk.handle(dict(cmd='model', ops=ops))['model']
            s = next(x for x in node(m1, 2200203, 11)[0]['slots'] if x['name'] == 'Count')
            self.assertEqual((s['source'], s['value'], s['definedIn'], s['own']), ('fixed', 2, 2200203, True))
            m2 = bk.handle(dict(cmd='model', ops=ops + [dict(op='applyToParent', bundle=2200203, node=11, slot='Count')]))['model']
            s1 = next(x for x in node(m2, 2200201, 11)[0]['slots'] if x['name'] == 'Count')
            s3 = next(x for x in node(m2, 2200203, 11)[0]['slots'] if x['name'] == 'Count')
            self.assertEqual((s1['definedIn'], s3['definedIn'], s3['own']), (2200201, 2200201, False))
            self.assertEqual(node(m2, 2200201, 11)[1]['usedBy'], [2100201, 2100202, 2100203], 'impact list')
            m3 = bk.handle(dict(cmd='model', ops=[dict(op='applyValueToParent', values=2300202, name='Bolt.Damage')]))['model']
            v = next(v for v in m3['values'] if v['key'] == 2300201)
            self.assertEqual(next(x['value'] for x in v['entries'] if x['name'] == 'Bolt.Damage'), 30)

    def test_new_skill_duplicate_and_keys(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            ops = [dict(op='newSkill', kind='Skill', ownerNumber=2, name='Test skill', moduleParent=2200001),
                   dict(op='duplicateSkill', owner=2100205, withEffects=True, name='Charge Beam copy')]
            m = bk.handle(dict(cmd='model', ops=ops))['model']
            k1, k2 = m['ops']
            self.assertEqual(len({k1['key'] % 100, k1['bundle'] % 100, k1['values'] % 100}), 1, 'skill, module and value keys share the serial')
            self.assertEqual(str(k2['key'])[:5], '21002')
            self.assertEqual(len(k2['effects']), 2, 'both beams copied')
            dup = next(o for o in m['owners'] if o['key'] == k2['key'])
            self.assertTrue(all(int(v) in dup['bundles'] for v in k2['effects'].values()), 'spawn nodes point at the copies')
            self.assertEqual(errors(m), [])
            self.assertEqual(k1['seeded'], ['Cast.Duration'], 'a value bundle without parent gets the names its bundle reads')

    def test_new_tier_skill_seeds_values_but_not_keyframes(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            ops = [dict(op='newSkill', kind='Skill', ownerNumber=1, name='Slash II', moduleParent=2200101, valueParent=None)]
            m = bk.handle(dict(cmd='model', ops=ops))['model']
            seeded = m['ops'][0]['seeded']
            self.assertIn('Cast.Duration', seeded)
            self.assertFalse([n for n in seeded if n.startswith('@')], 'keyframe references are not value names')

    def test_inserted_and_deleted_rows_keep_bundle_order(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            first = se.Loaded(se.Workspace(t.cfg))
            ops = [dict(op='setLink', bundle=2200101, node=11, slot='Count', source='Fixed', value=1),
                   dict(op='setLink', bundle=2200101, node=12, slot='PresentationKey', source=None)]
            self.assertTrue(bk.handle(dict(cmd='save', ops=ops, base=first.base, allowErrors=True))['ok'])
            wb = openpyxl.load_workbook(t.cfg.data_dir / 'ModuleBundles.xlsx')
            vals = [r[:5] for r in wb['Value Links'].iter_rows(min_row=8, values_only=True)]
            orig = [tuple(r[c] for c in ('BundleKey', 'NodeId', 'Slot', 'Source', 'Value')) for r in first.tables['ValueLinkData']['rows']]
            expected = [r for r in orig if r[:3] != (2200101, 12, 'PresentationKey')]
            at = max(i for i, r in enumerate(expected) if r[0] == 2200101) + 1
            expected.insert(at, (2200101, 11, 'Count', 'Fixed', 1))
            self.assertEqual(vals, expected, 'new row under the bundle, deleted row pulls the rest up, other rows untouched')
            f, ref = wb['Value Links'].cell(8 + at, 1).font, wb['Value Links'].cell(8, 1).font
            self.assertEqual((f.name, f.sz, f.b), (ref.name, ref.sz, ref.b), 'new row copies the style above')

    def test_save_refuses_errors(self):
        with Temp() as t:
            res = se.Backend(t.cfg).handle(dict(cmd='save', ops=[dict(op='setLink', bundle=2200001, node=1, slot='Duration', source=None)]))
            self.assertFalse(res['ok'])

    def test_lock_and_external_change_block_save(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            first = se.Loaded(se.Workspace(t.cfg))
            op = [dict(op='setValue', values=2300201, name='Bolt.Damage', value=26)]
            owner = t.cfg.data_dir / '~$Skills.xlsx'
            owner.write_bytes(b'x')
            res = bk.handle(dict(cmd='save', ops=op, base=first.base))
            self.assertFalse(res['ok'])
            self.assertIn('Excel has Skills.xlsx open', res['blocked'][0])
            owner.unlink()
            res = bk.handle(dict(cmd='save', ops=op, base=dict(first.base, Skills='0' * 64)))
            self.assertFalse(res['ok'])
            self.assertIn('changed after the editor read it', res['blocked'][0])
            self.assertEqual(sha(t.cfg.data_dir / 'Skills.xlsx'), first.base['Skills'], 'a blocked save changes nothing')

    def test_practice_space_is_isolated(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            product_before = {p.name: sha(p) for p in t.cfg.data_dir.glob('*.xlsx')}
            logic_before = build(t.cfg)['manifest']['logicSha256']
            self.assertTrue(bk.handle(dict(cmd='practice', action='create'))['created'])
            lo = bk.loaded('practice')
            ops = [dict(op='newSkill', kind='Skill', name='Practice skill', moduleParent=2200001)]
            m = bk.handle(dict(cmd='model', workspace='practice', ops=ops))['model']
            self.assertEqual(str(m['ops'][0]['key'])[3:5], '98', 'practice owner number')
            self.assertEqual(errors(m), [])
            self.assertTrue(bk.handle(dict(cmd='save', workspace='practice', ops=ops, base=lo.base))['ok'])
            self.assertEqual({p.name: sha(p) for p in t.cfg.data_dir.glob('*.xlsx')}, product_before)
            self.assertEqual(build(t.cfg)['manifest']['logicSha256'], logic_before)
            from skillgraph.pipeline import run
            r = run(se.Workspace(t.cfg, 'practice').files, t.cfg, True)
            self.assertIn('PracticeKey', [x.code for x in r['errors']], 'practice keys are errors in product mode')

    def test_aux_memo_combo_bands(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)
            first = se.Loaded(se.Workspace(t.cfg))
            ops = [dict(op='setMemo', bundle=2200002, node=3, text='shared damage'),
                   dict(op='saveCombo', name='Damage pack', bundle=2200002, nodes=[3]),
                   dict(op='insertCombo', bundle=2200214, combo='Damage pack', event='Hit'),
                   dict(op='setBands', bundle=2200201, bands=[dict(name='Wind-up', **{'from': 'CastStart', 'to': 'CastStart'})]),
                   dict(op='setCollapsed', bundle=2200001, collapsed=True)]
            m = bk.handle(dict(cmd='model', ops=ops))['model']
            self.assertEqual(m['aux']['bundles']['2200002']['memos']['3'], 'shared damage')
            n, _ = node(m, 2200214, m['ops'][2]['nodes'][0])
            self.assertEqual((n['module'], n['event']), ('Damage', 'Hit'))
            res = bk.handle(dict(cmd='save', ops=ops[:2] + ops[3:], base=first.base))
            self.assertTrue(res['ok'], res)
            aux = json.loads(t.cfg.aux_path(t.cfg.data_dir).read_text(encoding='utf-8'))
            self.assertEqual((aux['format'], aux['combos'][0]['name'], aux['bundles']['2200001']['collapsed']), (t.cfg['formats']['aux'], 'Damage pack', True))
            self.assertEqual(res['changes'], [], 'aux-only edits do not touch the workbooks')

    def test_columns_add_change_remove(self):
        with Temp() as t:
            bk = se.Backend(t.cfg)

            def model(ops):
                r = bk.handle(dict(cmd='model', ops=ops))
                self.assertTrue(r['ok'], r.get('error'))
                return r['model']

            def fail(ops):
                with self.assertRaises(se.OpError) as cm:
                    bk.handle(dict(cmd='model', ops=ops))
                return str(cm.exception)

            add = dict(op='addColumn', bundle=2200201, event='Moment', eventValue=0.5)
            m = model([add])
            self.assertEqual(m['aux']['bundles']['2200201']['columns'], [{'event': 'Moment', 'value': 0.5, 'from': None}])
            self.assertIn('already exists', fail([add, add]))
            self.assertIn('already exists', fail([dict(op='addColumn', bundle=2200201, event='Moment', eventValue='Bolt.Release')]))
            self.assertIn('not available', fail([dict(op='addColumn', bundle=2200201, event='Hit')]))
            self.assertIn('already exists', fail([add, dict(op='addColumn', bundle=2200202, event='Moment', eventValue=0.5)]), 'parent layer column')
            m = model([add, dict(op='addNode', bundle=2200201, event='Moment', eventValue=0.5, module='PlaySound', line='Presentation')])
            self.assertNotIn('columns', m['aux']['bundles'].get('2200201', {}), 'a column that receives a node leaves the aux file')

            m = model([dict(op='setColumn', bundle=2200201, event='Moment', eventValue='Bolt.Release', toEvent='Moment', toEventValue=0.35)])
            self.assertEqual({node(m, 2200201, i)[0]['eventValue'] for i in (11, 12)}, {0.35})
            self.assertEqual(node(m, 2200202, 11)[0]['eventValue'], 0.35, 'tier 2 inherits')
            m = model([dict(op='setColumn', bundle=2200202, event='Moment', eventValue='Bolt.Release', toEvent='Moment', toEventValue=0.4)])
            self.assertEqual((node(m, 2200202, 11)[0]['eventValue'], node(m, 2200201, 11)[0]['eventValue']), (0.4, 'Bolt.Release'), 'child layer overrides time only')
            self.assertIn('time only', fail([dict(op='setColumn', bundle=2200202, event='Moment', eventValue='Bolt.Release', toEvent='CastStart')]))

            m = model([dict(op='removeColumn', bundle=2200202, event='Moment', eventValue='Bolt.Release')])
            self.assertEqual(sorted(r['id'] for r in node(m, 2200202, 1)[1]['removed']), [11, 12], "inherited nodes get '-'")
            self.assertIsNotNone(node(m, 2200201, 11)[0], 'parent layer untouched')
            m = model([dict(op='removeColumn', bundle=2200205, event='CheckPass', fromNode=10)])
            self.assertEqual(sorted(m['ops'][0]['removed']), [11, 13])
            m = model([dict(op='removeColumn', bundle=2200205, event='CastStart')])
            self.assertEqual(sorted(m['ops'][0]['removed']), [1, 10, 11, 12, 13], "inherited cast lock gets '-'; exit columns of the removed check go too")

    def test_bad_op_reports_position(self):
        with Temp() as t:
            with self.assertRaisesRegex(se.OpError, 'Operation #2'):
                se.apply_ops(se.Loaded(se.Workspace(t.cfg)), [dict(op='setMemo', bundle=1, node=1, text='a'),
                                                              dict(op='setLink', bundle=2999999, node=1, slot='X', source='Ref', value='A.B')])

    def test_serve(self):
        with Temp() as t:
            req = [dict(id=1, cmd='ping'), dict(id=2, cmd='model', ops=[dict(op='setMemo', bundle=2200001, node=1, text='hi')]),
                   dict(id=3, cmd='nope'), 'not json']
            stdin = io.BytesIO(b''.join((json.dumps(r) if isinstance(r, dict) else r).encode() + b'\n' for r in req))
            out = io.BytesIO()
            se.serve(t.cfg, stdin, out)
            res = [json.loads(x) for x in out.getvalue().decode('utf-8').splitlines()]
            self.assertEqual([(r['id'], r['ok']) for r in res], [(1, True), (2, True), (3, False), (None, False)])
            self.assertEqual(res[1]['model']['aux']['bundles']['2200001']['memos']['1'], 'hi')


if __name__ == '__main__':
    unittest.main()
