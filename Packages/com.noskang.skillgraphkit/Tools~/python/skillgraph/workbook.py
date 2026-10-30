"""Excel side: header contract, typed cell parsing, reading tables, lock check and atomic save.

Header (7 rows by default): 1 JSON name, 2 description, 3 type, 4 column name, 5 export token, 6 reference, 7 range.
Data starts on the row after the header. Sheets whose name starts with "!" are notes and are never read as data.
The literal cell text "<empty>" stores an empty string (a blank cell means "use the column default").
"""
import os
import re
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .core import LOC_RE, NUMERIC, Report, SkillDataError, is_num
from .messages import text

EMPTY = '<empty>'


def parse(c, raw, rep, at, enums, cfg):
    kind = c['type']
    if kind in ('string', 'loc'):
        if not isinstance(raw, str):
            rep.e('TextRequired', at, raw=raw); return None
        value = raw.strip()
        if kind == 'loc' and not LOC_RE.match(value):
            rep.e('LocFormat', at, raw=raw); return None
        return value
    if kind in NUMERIC or kind.startswith('key:'):
        if not is_num(raw):
            rep.e('NumberRequired', at, raw=raw); return None
        if (kind == 'int' or kind.startswith('key:')) and raw != int(raw):
            rep.e('IntegerRequired', at, raw=raw); return None
        if kind.startswith('key:'):
            band = cfg.band(kind[4:])
            if not cfg.is_key(int(raw), band):
                rep.e('KeyBand', at, band=band, raw=raw); return None
        return int(raw) if (kind == 'int' or kind.startswith('key:')) else float(raw)
    if kind.startswith('enum:'):
        name = kind.split(':')[1]
        if not isinstance(raw, str):
            rep.e('EnumName', at, raw=raw); return None
        if raw not in enums[name]:
            rep.e('InvalidEnum', at, raw=raw, allowed=', '.join(enums[name])); return None
        return raw
    if kind.startswith('list:'):
        sub = dict(c, type=kind[5:])
        items = [x.strip() for x in str(raw).split(',')]
        if any(x == '' for x in items):
            rep.e('EmptyListItem', at, raw=raw); return None
        return [parse(sub, int(x) if re.fullmatch(r'-?\d+', x) else x, rep, at, enums, cfg) for x in items]
    if kind == 'value':
        if isinstance(raw, bool) or not (is_num(raw) or isinstance(raw, str)):
            rep.e('ValueType', at, raw=raw); return None
        return raw.strip() if isinstance(raw, str) else raw
    raise ValueError(kind)


def token(reg, export):
    return reg['tokens'].get(export, export)


def read_table(wb, t, rep, reg, cfg):
    if t['sheet'] not in wb.sheetnames:
        rep.e('MissingSheet', dict(doc=t['document']), sheet=t['sheet']); return []
    ws, cols, h = wb[t['sheet']], t['columns'], t['headerRows']
    sheet = dict(sheet=t['sheet'])
    if ws.sheet_state != 'visible':
        rep.e('HiddenDataSheet', sheet)
    if ws.merged_cells.ranges:
        rep.e('MergedCells', sheet)
    if ws['A1'].value != t['json']:
        rep.e('HeaderJsonName', dict(sheet=t['sheet'], col='A1'), got=ws['A1'].value, want=t['json'])
    checks = [(3, lambda c: c['type']), (4, lambda c: c['name'])]
    if h >= 5:
        checks.append((5, lambda c: token(reg, c['export'])))
    for r, pick in checks:
        got = [ws.cell(r, i + 1).value for i in range(len(cols))]
        want = [pick(c) for c in cols]
        if got != want:
            rep.e('HeaderMismatch', dict(sheet=t['sheet'], row=r), got=got, want=want)
    if h >= 5:
        row5 = [ws.cell(5, i + 1).value for i in range(len(cols))]
        if any(v not in reg['tokens'].values() for v in row5):
            rep.e('HeaderRowsGuess', dict(sheet=t['sheet'], row=5), tokens=', '.join(reg['tokens'].values()))
    rows = []
    for r in range(h + 1, ws.max_row + 1):
        cells = [ws.cell(r, i + 1) for i in range(len(cols))]
        if all(c.value is None for c in cells):
            continue
        extra = [ws.cell(r, i).value for i in range(len(cols) + 1, ws.max_column + 1)]
        if any(v is not None for v in extra):
            rep.e('ExtraCell', dict(sheet=t['sheet'], row=r))
        row = {'_row': r, '_sheet': t['sheet']}
        for c, cell in zip(cols, cells):
            at = dict(sheet=t['sheet'], row=r, col=c['name'])
            if cell.data_type == 'f':
                rep.e('FormulaNotAllowed', at); row[c['name']] = None; continue
            raw = cell.value
            if raw is None or raw == '':
                if c['required']:
                    rep.e('RequiredCell', at)
                row[c['name']] = c['default']
                continue
            if raw == EMPTY:
                row[c['name']] = ''; continue
            row[c['name']] = parse(c, raw, rep, at, reg['enums'], cfg)
        rows.append(row)
    seen = set()
    for row in rows:
        k = tuple(row.get(x) for x in t['key'])
        if k in seen:
            rep.e('DuplicateKey', dict(sheet=t['sheet'], row=row['_row']), key=k)
        seen.add(k)
    return rows


def open_workbook(path):
    return openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=False)


def load(files, rep, reg, cfg):
    data = {t['json']: [] for t in reg['tables']}
    for doc, name in files.items():
        if name is None:
            continue
        wb = open_workbook(name)
        try:
            known = {t['sheet'] for t in reg['tables'] if t['document'] == doc}
            for s in wb.sheetnames:
                if s not in known and not s.startswith('!'):
                    rep.e('SheetsMismatch', dict(doc=doc), sheet=s)
            for t in reg['tables']:
                if t['document'] == doc:
                    data[t['json']] = read_table(wb, t, rep, reg, cfg)
        finally:
            wb.close()
    return data


# ------------------------------------------------------------------ writing
def owner_file(path):
    path = Path(path)
    return path.with_name('~$' + path.name)


def lock_state(path, lang):
    """None when the file can be replaced; otherwise the reason (Excel owner file, or the file is open elsewhere)."""
    path = Path(path)
    owner = owner_file(path)
    if owner.exists():
        return text(lang, 'save.owner', file=path.name, owner=owner.name)
    if path.exists():
        try:
            with open(path, 'r+b'):
                pass
        except OSError as error:
            return text(lang, 'save.locked', file=path.name, error=error)
    return None


def lock_check(path, lang):
    """For generated views: an owner file is only a warning; a locked file is an error."""
    path = Path(path)
    warnings = []
    if owner_file(path).exists():
        warnings.append(text(lang, 'lock.owner', owner=owner_file(path).name))
    if path.exists():
        try:
            with open(path, 'r+b'):
                pass
        except OSError as error:
            raise SkillDataError(text(lang, 'lock.locked', path=path, error=error))
    return warnings


def atomic_save(wb, path, backup=None):
    """Save to a temp file in the same folder, keep one previous copy, then replace."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix='.' + path.stem + '.', suffix=path.suffix, dir=path.parent)
    os.close(fd)
    try:
        wb.save(name)
        if path.exists():
            backup = Path(backup) if backup else path.with_name(path.stem + '.previous' + path.suffix)
            backup.parent.mkdir(parents=True, exist_ok=True)
            if backup.exists():
                backup.unlink()
            os.replace(path, backup)
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


HEAD = PatternFill('solid', fgColor='173D50')
SUB = PatternFill('solid', fgColor='E8EEF2')


def write_header(ws, t, reg):
    """Write the header rows of a registry table onto an empty sheet."""
    cols = t['columns']
    ws['A1'] = t['json']
    ws['A1'].font = Font(bold=True, size=12)
    rows = [[c.get('desc') or '' for c in cols], [c['type'] for c in cols], [c['name'] for c in cols],
            [token(reg, c['export']) for c in cols], [c.get('ref') or '' for c in cols], [c.get('range') or '' for c in cols]]
    for r, values in enumerate(rows, start=2):
        for i, v in enumerate(values, start=1):
            cell = ws.cell(r, i, v if v != '' else None)
            cell.font = Font(bold=(r == 4), color='FFFFFF' if r == 4 else '333333', size=10)
            cell.fill = HEAD if r == 4 else SUB
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = ws.cell(t['headerRows'] + 1, 1)


def new_report(cfg):
    return Report(cfg['language'])
